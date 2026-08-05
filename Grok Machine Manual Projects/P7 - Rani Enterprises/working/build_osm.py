#!/usr/bin/env python3
"""Build Robotic Sealant Dispenser OSM — Rani Enterprises (python-docx).
Writes only under Grok Machine Manual Projects. Navy/Calibri format.
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "RaniEnterprises_RoboticSealantDispenser_OSM_V1.0.docx"
DATA = json.loads((ROOT / "inputs" / "machine_data.json").read_text(encoding="utf-8"))

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x55, 0x55, 0x55)
WARN_BG = "FFF2CC"
NOTICE_BG = "F2F2F2"
ALT_ROW = "D9E1F2"
PLACE_BG = "F7F7F7"
OA_BG = "FCE4D6"


def set_run(run, *, size=11, bold=False, color=BLACK, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, color="1F3864", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_para(cell, text, *, bold=False, size=10, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)
    return p


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run(run, size=9, color=NAVY)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_end)


def configure_section(section, *, different_first=False):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.header_distance = Mm(12)
    section.footer_distance = Mm(12)
    section.different_first_page_header_footer = different_first


def build_header_footer(doc: Document):
    section = doc.sections[0]
    configure_section(section, different_first=True)
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""

    header = section.header
    hp = header.paragraphs[0]
    hp.clear()
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = hp.add_run("Robotic Sealant Dispenser Machine")
    set_run(r1, size=9, color=NAVY)
    r2 = hp.add_run("  |  Operating & Safety Manual  |  Page ")
    set_run(r2, size=9, color=NAVY)
    add_page_number(hp)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("CONFIDENTIAL — Draft  |  Rani Enterprises")
    set_run(fr, size=9, bold=True, color=GREY)


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.page_break_before = True
    r = p.add_run(text)
    set_run(r, size=14, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    pPr.append(outline)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run(r, size=12, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "1")
    pPr.append(outline)
    return p


def body(doc, text):
    p = None
    for para in text.strip().split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(para.strip() if para.strip() else "")
        set_run(r, size=11, color=BLACK)
    return p


def bullet(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        r = p.add_run(item)
        set_run(r, size=11)
        p.paragraph_format.space_after = Pt(3)


def alert_box(doc, kind: str, lines: list[str]):
    colors = {
        "WARNING": (WARN_BG, NAVY),
        "NOTICE": (NOTICE_BG, BLACK),
        "PLACEHOLDER": (PLACE_BG, BLACK),
        "OPEN ACTION": (OA_BG, RGBColor(0xCC, 0, 0)),
        "CAUTION": (OA_BG, RGBColor(0xC6, 0x5A, 0x00)),
    }
    bg, title_color = colors.get(kind, (NOTICE_BG, BLACK))
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, bg)
    set_cell_border(cell, "000000" if kind == "WARNING" else "AAAAAA")
    cell.text = ""
    p0 = cell.paragraphs[0]
    r0 = p0.add_run(kind)
    set_run(r0, size=11, bold=True, color=title_color)
    for line in lines:
        p = cell.add_paragraph()
        r = p.add_run(line)
        set_run(r, size=10, bold=(kind == "PLACEHOLDER"), color=BLACK)
    doc.add_paragraph()


def styled_table(doc, headers, rows, col_widths_cm=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell_para(cell, h, bold=True, size=10, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(cell, "1F3864")
        set_cell_border(cell)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell_para(cell, str(val), size=9)
            set_cell_border(cell)
            if ri % 2 == 1:
                shade_cell(cell, ALT_ROW)
    if col_widths_cm:
        for row in table.rows:
            for i, w in enumerate(col_widths_cm):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def flowchart(doc, title: str, steps: list[str]):
    if title:
        h2(doc, title)
    table = doc.add_table(rows=len(steps), cols=1)
    for i, step in enumerate(steps):
        cell = table.rows[i].cells[0]
        is_end = i == 0 or i == len(steps) - 1
        if is_end:
            shade_cell(cell, "1F3864")
            cell_para(cell, step, bold=True, size=11, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            shade_cell(cell, "FFFFFF")
            cell_para(cell, step, bold=False, size=11, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_border(cell, "1F3864", "12")
    doc.add_paragraph()


def cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("OPERATING & SAFETY MANUAL")
    set_run(r, size=22, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Robotic Sealant Dispenser Machine")
    set_run(r, size=18, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Rani Enterprises")
    set_run(r, size=14, color=NAVY)

    doc.add_paragraph()
    meta = [
        ("Manufacturer (OEM)", DATA["manufacturer"]),
        ("End User", DATA["end_user"]),
        ("Machine Type", DATA["machine_type"]),
        ("Document Title", "Operating & Safety Manual"),
        ("Document No.", "RANI-SEALANT-OSM-001"),
        ("Revision", DATA["revision"]),
        ("Date", DATA["date"]),
        ("Prepared By", "Certified Safety Professional"),
        ("Status", DATA["status"]),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    for i, (k, v) in enumerate(meta):
        cell_para(table.rows[i].cells[0], k, bold=True, size=11, color=NAVY)
        cell_para(table.rows[i].cells[1], v, size=11)
        set_cell_border(table.rows[i].cells[0])
        set_cell_border(table.rows[i].cells[1])
        shade_cell(table.rows[i].cells[0], ALT_ROW)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CONFIDENTIAL — Draft for review")
    set_run(r, size=9, bold=True, color=GREY)


def revision_page(doc):
    h1(doc, "Revision History")
    body(doc, "Four-column revision table (Prepared By column not used).")
    styled_table(
        doc,
        ["Revision", "Date", "Description", "Reviewed By"],
        [
            ["V1.0", DATA["date"], "Initial draft from site photos and intake notes", "Certified Safety Professional"],
            ["V2.0", "TBD", "Post customer review / closed open actions", "OEM / End User"],
        ],
        [2.5, 3, 7.5, 4],
    )


def toc_page(doc):
    h1(doc, "Table of Contents")
    body(
        doc,
        "H1 entries only. After opening in Word: right-click → Update Field → Update entire table if a TOC field is inserted later. This edition uses a static TOC list.",
    )
    for e in DATA["section_map"]:
        p = doc.add_paragraph()
        r = p.add_run(e)
        set_run(r, size=11, color=BLACK)
        p.paragraph_format.space_after = Pt(4)


def section_1(doc):
    h1(doc, "1. Front Matter")
    h2(doc, "1.1 Document Scope")
    body(
        doc,
        """This Operating & Safety Manual (OSM) provides information for the safe operation, energy isolation, and basic maintenance of the Robotic Sealant Dispenser Machine manufactured by Rani Enterprises.

The machine integrates a Fisnar F4303N ADVANCE robot/dispenser system with a Mitsubishi PLC/HMI cell controller, automatic front door, side-door interlocks, two-hand cycle start, and pneumatic services.

This document is prepared from site photographs and intake notes. A formal Machine Safety Assessment was not available at the time of drafting. Items marked PLACEHOLDER or Open Action must be closed before treating the manual as customer-released.

Preparation follows the principles of ISO 20607:2019, ISO 12100:2010, IEC 82079-1:2019, and IS 4571:2008. No performance-level (PLr) claims are made without an independent assessment.""",
    )
    h2(doc, "1.2 Intended Audience")
    bullet(
        doc,
        [
            "Production operators (instructed persons) who load components, start cycles, and unload parts",
            "Maintenance and troubleshooting personnel (skilled persons)",
            "Supervisors and safety representatives",
            "Contractors working near the machine (awareness only)",
        ],
    )
    body(doc, "Untrained persons shall not operate, teach, adjust, or maintain this machine.")

    h2(doc, "1.3 General Safety Rules")
    bullet(
        doc,
        [
            "Read this manual before operating or servicing the machine.",
            "Only trained and authorised persons may operate or maintain the machine.",
            "Never bypass two-hand controls, door interlocks, E-Stops, or guards.",
            "Apply full Lockout/Tagout (Section 4) before any intervention inside the process zone, panel work, or pneumatic work.",
            "Keep hands clear of the robot path, dispenser tip, and automatic front door during motion.",
            "Do not reach into the cell while the robot is running or the front door is closing.",
            "Use parking and purging functions as required to prevent tip dry-out and unplanned interventions.",
            "Report damaged safety devices immediately and take the machine out of service until restored.",
            "Wear PPE required by the end-user site and the sealant Safety Data Sheet (SDS).",
        ],
    )

    h2(doc, "1.4 Signal Word Hierarchy")
    styled_table(
        doc,
        ["Signal Word", "Meaning"],
        [
            ["WARNING", "Hazardous situation that could result in death or serious injury if not avoided."],
            ["CAUTION", "Hazardous situation that could result in minor or moderate injury if not avoided."],
            ["NOTICE", "Practices not related to personal injury, or important operational information."],
            ["PLACEHOLDER", "Data not available from source inputs; must be confirmed."],
            ["OPEN ACTION", "Tracked gap — see standalone Open Actions register."],
        ],
        [4, 12],
    )
    alert_box(doc, "NOTICE", ["DANGER signal-word boxes are not used in this manual."])

    h2(doc, "1.5 Prepared By")
    body(doc, "Prepared by: Certified Safety Professional\nRole only — no individual name is used in customer-facing attribution fields of this manual.")


def section_2(doc):
    h1(doc, "2. Machine Overview")
    h2(doc, "2.1 Machine Identification")
    styled_table(
        doc,
        ["Field", "Value"],
        [
            ["Machine name", DATA["machine_name"]],
            ["Manufacturer (OEM)", DATA["manufacturer"]],
            ["End user", DATA["end_user"]],
            ["Machine type", DATA["machine_type"]],
            ["Integrated robot", DATA["robot"]["make_model"]],
            ["Label observed", "ROBOTIC SEALANT DISPENSER MACHINE"],
            ["Model / Serial / Year", "[PLACEHOLDER — confirm from nameplate — OA-P7-002]"],
            ["Site address", DATA["site"]],
        ],
        [5, 11],
    )
    alert_box(
        doc,
        "PLACEHOLDER",
        [
            "[PLACEHOLDER — Model No., Serial No., Asset No., Year of manufacture, power rating — OA-P7-002]",
            "Photo reference: Images/Nameplate.jpg (title label only; no electrical nameplate data).",
        ],
    )

    h2(doc, "2.2 Technical Specifications")
    e, pneu = DATA["electrical"], DATA["pneumatic"]
    hmi_s = DATA["hmi"]["settings_observed"]
    styled_table(
        doc,
        ["Parameter", "Value"],
        [
            ["Electrical supply voltage", e["voltage"]],
            ["Phases / frequency", f"{e['phases']} / {e['frequency']}"],
            ["Machine power rating", e["power_rating"]],
            ["Control power", e["control_voltage"]],
            ["PLC", e["plc"]],
            ["HMI", e["hmi"]],
            ["PSU observed", e["psu"]],
            ["Main isolator", e["isolator"]],
            ["Robot", DATA["robot"]["make_model"]],
            ["Plant air (nominal)", pneu["plant_pressure_nominal"]],
            ["Dispenser inlet set range (HMI)", pneu["dispenser_inlet_setpoint_range"]],
            ["Hydraulic", "Not applicable"],
            ["Operating modes", ", ".join(DATA["operating_modes"])],
            ["Models observed on HMI", ", ".join(DATA["models_supported"])],
            ["Cycle time (specified)", DATA["cycle_time"]["specified"]],
            ["Cycle time (observed instance)", DATA["cycle_time"]["observed_instance"]],
            ["Idle time to parking (HMI)", f"{hmi_s['idle_time_to_go_parking_min']} minutes"],
            ["Sealant capacity", DATA["capacity"]],
            ["Axis / guide speed", DATA["robot"]["axis_speed"]],
        ],
        [6, 10],
    )

    h2(doc, "2.3 Intended Use")
    body(doc, DATA["process_summary"])
    body(
        doc,
        """Intended users: trained production operators and skilled maintenance personnel.

Reasonably foreseeable misuse includes: reaching into the cell during robot motion or door closure; bypassing two-hand, door, or E-Stop devices; teaching or changing robot programs without authorisation; operating with side doors open or interlocks defeated; running without adequate air pressure; ignoring parking/purge requirements; and performing electrical/pneumatic work without LOTO.""",
    )

    h2(doc, "2.4 Layout & Key Zones")
    bullet(
        doc,
        [
            "Operator load/unload face with two-hand stations and automatic front door",
            "Process zone: Fisnar robot, dispenser, fixture tray, parking and purging cups",
            "Left/right side doors with non-contact safety switches",
            "Upper front control face: HMI and indicators",
            "Side control panel (Bhartia BCH) with isolator, FRL, PLC/PSU",
            "Tower light on top structure",
        ],
    )
    alert_box(
        doc,
        "NOTICE",
        [
            "Image placeholders (customer or later embed): Images/Full view Front.jpg; Full view side.jpg; Component on tray.jpg.",
        ],
    )

    h2(doc, "2.5 Key Components")
    bullet(
        doc,
        [
            "Fisnar F4303N ADVANCE cartesian robot base and head",
            "Fisnar blue sealant dispenser cartridge/syringe assembly",
            "Component fixture tray / nest",
            "Parking cup and purging cup for dispenser tip",
            "Automatic front safety door (pneumatically assisted)",
            "Left and right side access doors with Omron GLS-M1 non-contact switch",
            "Dual two-hand cycle-start stations (left RESET + green; right green + E-Stop)",
            "Mitsubishi HMI and PLC cell control",
            "Mean Well 24 VDC power supply (NDR-240-24 observed)",
            "Lauritz Knudsen rotary main isolator (POWER ON/OFF)",
            "SMC FRL with pressure gauge / pressure-switch interface",
            "Status tower light",
        ],
    )


def section_3(doc):
    h1(doc, "3. Hazard Identification")
    body(
        doc,
        """This section is a narrative summary of hazards reasonably identified from machine construction, photographs, HMI status/alarms, and intake process description. It is not a substitute for a formal Machine Safety Assessment.""",
    )
    alert_box(
        doc,
        "OPEN ACTION",
        [
            "OA-P7-001: No formal Machine Safety Assessment was available for this machine.",
            "An Independent Safety Consultant should perform a Machine Safety Assessment before claiming compliance completeness or performance levels.",
        ],
    )
    body(
        doc,
        """Principal hazards associated with this equipment include:

Mechanical / robot motion. The Fisnar robot and dispenser head move within the enclosure. Contact with the moving head, tip, or cable track can cause impact, puncture, or shearing injury. Labels on the robot structure warn operators to keep hands clear of moving parts.

Automatic front door. The front door closes from the top during the cycle (intake description). Crush or impact injury can occur if hands or body are in the doorway during closure. The cell does not include a light curtain at the load face (intake). Residual risk depends on door design, two-hand start discipline, and correct interlock function (OA-P7-005).

Side door access. Left and right side doors provide maintenance access. Omron non-contact switches and HMI side-door status are present. Opening a side door during motion, or defeating the sensor, can expose persons to robot motion.

Pneumatic energy. Plant air (about 6 bar at the FRL photo) drives door and process functions. Unexpected motion or residual pressure after shutdown can injure during maintenance if air is not isolated and dumped.

Electrical energy. The control panel contains mains distribution, a 24 VDC supply, PLC, and relays. Shock and arc hazards exist during live troubleshooting. Electrical work requires isolation at the main POWER ON/OFF switch and verification by a skilled person.

Chemical / sealant. Sealant materials may present skin, eye, or inhalation hazards and may solidify in the tip if idle too long. Product identity and SDS were not provided (OA-P7-006). Parking and purging reduce process failure and secondary unsafe interventions.

Control system faults / low air. HMI live alarms observed include Air Pressure Low, Robot Emergency Pressed [X12], Front Auto Door open Error [X6], and Machine Emergency Pressed [X0]. Operators must treat alarms as stop conditions and clear causes before restart.

Administrative residual risk. Until a Machine Safety Assessment is completed, residual risk ratings and required additional protective measures remain incomplete. This manual uses PLACEHOLDER and Open Actions rather than inventing scores or PLr values.

End of Section 3. Commission a Machine Safety Assessment and update this manual when the report is available.""",
    )
    alert_box(
        doc,
        "NOTICE",
        [
            "Open Actions are maintained as a standalone Excel file — not inside this manual.",
            "File: RaniEnterprises_RoboticSealantDispenser_OpenActions_V1.0.xlsx",
        ],
    )


def section_4(doc):
    h1(doc, "4. Lockout / Tagout (LOTO) Procedures")
    h2(doc, "4.1 Energy Isolation Points")
    body(
        doc,
        """Isolate all hazardous energy before entering the process zone for cleaning near the robot path, tip change, fixture work, panel access, or any task where unexpected motion or energisation could cause injury.

Energy sources: Electrical — supply via main rotary isolator (voltage PLACEHOLDER — OA-P7-002); Pneumatic — plant air via FRL / supply isolation (nominal ~6 bar); Hydraulic — not applicable.""",
    )
    rows = [
        [lp["id"], lp["energy_type"], lp["image_code"], lp["isolation_method"], lp["location"]]
        for lp in DATA["loto_points"]
    ]
    styled_table(
        doc,
        ["Point ID", "Energy Type", "Image Code", "Isolation Method", "Location"],
        rows,
        [2, 2.5, 2.5, 5.5, 4.5],
    )
    alert_box(
        doc,
        "PLACEHOLDER",
        ["[PLACEHOLDER — Confirm lockable air isolation valve and isolator hasp/padlock hardware — OA-P7-007]"],
    )

    h2(doc, "4.2 LOTO Equipment Required")
    bullet(
        doc,
        [
            "Personal padlocks (one key per authorised person)",
            "Lockout hasps for multi-person lockout",
            "Danger/lockout tags with name, date, reason",
            "Voltage tester / proving unit (skilled electrical person)",
            "[PLACEHOLDER — site-specific LOTO kit inventory]",
        ],
    )

    flowchart(
        doc,
        "4.3 Tag-Out Sequence",
        [
            "START",
            "Notify affected personnel",
            "Stop machine in safe state",
            "Turn isolator to OFF",
            "Isolate plant air supply",
            "Dump residual air pressure",
            "Apply locks and tags",
            "Verify zero energy",
            "Perform work",
            "END",
        ],
    )
    flowchart(
        doc,
        "4.4 Tag-In Sequence",
        [
            "START",
            "Confirm tools removed",
            "Confirm guards doors closed",
            "Remove personal locks tags",
            "Restore air supply",
            "Turn isolator to ON",
            "Reset E-Stops if needed",
            "Function-check safety devices",
            "Return to production",
            "END",
        ],
    )


def section_5(doc):
    h1(doc, "5. Emergency Stop")
    h2(doc, "5.1 E-Stop Devices")
    body(
        doc,
        """Two emergency-stop actuators are photographically confirmed. E-Stop is for emergency conditions. It is not a substitute for LOTO during maintenance.""",
    )
    styled_table(
        doc,
        ["ID", "Location", "Type", "Source"],
        [[es["id"], es["location"], es["type"], es.get("source", "Photo")] for es in DATA["e_stops"]],
        [2, 6.5, 3.5, 4],
    )
    alert_box(
        doc,
        "NOTICE",
        [
            "ES-01: right two-hand station E STOP — HMI has shown Machine Emergency Pressed [X0].",
            "ES-02: Fisnar controller EMERGENCY STOP — HMI has shown Robot Emergency Pressed [X12].",
        ],
    )

    flowchart(
        doc,
        "5.2 Activation Sequence",
        [
            "START",
            "Detect emergency condition",
            "Press nearest E-Stop hard",
            "Machine motion stops",
            "Keep clear of process zone",
            "Alert supervisor if needed",
            "Do not reset until safe",
            "END",
        ],
    )
    flowchart(
        doc,
        "5.3 Reset Sequence",
        [
            "START",
            "Identify why E-Stop pressed",
            "Clear hazard and personnel",
            "Close side doors fully",
            "Twist-reset E-Stop device",
            "Clear HMI emergency alarm",
            "Press RESET if required",
            "Confirm ready status",
            "Restart only if safe",
            "END",
        ],
    )

    h2(doc, "5.4 After E-Stop")
    bullet(
        doc,
        [
            "Investigate cause (operator, robot fault, door fault, air pressure, etc.).",
            "Check HMI ALARM list and status indicators (EMERGENCY, ROBOT EMERGENCY).",
            "Confirm front automatic door and side-door statuses before resuming AUTO.",
            "If the dispenser tip was mid-cycle, follow parking/purge guidance before long idle.",
            "Repeated unexplained E-Stops → take out of service and call skilled maintenance.",
        ],
    )


def section_6(doc):
    h1(doc, "6. Operating Instructions")
    h2(doc, "6.1 Pre-Start Checks")
    bullet(
        doc,
        [
            "Guards and side doors closed and undamaged; Omron switches seated",
            "Front automatic door free of obstruction",
            "Main isolator ON; HMI powered; tower light functioning",
            "Plant air present; FRL gauge near expected operating pressure",
            "No active emergency on HMI; E-Stops released",
            "Correct model selected (e.g. JPTP / MGTP / APTDPT)",
            "Sealant cartridge installed and within process life; tip clean",
            "Work area clear; authorised operator only",
        ],
    )

    flowchart(
        doc,
        "6.2 Startup",
        [
            "START",
            "Turn POWER isolator ON",
            "Confirm air supply ready",
            "Release all E-Stops",
            "HMI boots to GLUE DISPENSER",
            "Select AUTO or MANUAL",
            "Select model program",
            "Clear any alarms",
            "Verify doors closed",
            "END",
        ],
    )

    flowchart(
        doc,
        "6.3 Automatic Production Cycle",
        [
            "START",
            "Select model on HMI",
            "Place component on tray",
            "Hands on both green starts",
            "Press two-hand together",
            "Front door closes",
            "Robot dispenses sealant",
            "Cycle ends door opens",
            "Unload finished part",
            "END",
        ],
    )
    alert_box(
        doc,
        "WARNING",
        [
            "Keep hands outside the door path and robot envelope during the cycle.",
            "Do not reach in if the door fails to close or the robot is still moving.",
        ],
    )

    h2(doc, "6.4 Model Selection")
    body(
        doc,
        """From HMI MODEL SELECTION: running model and robot program number are displayed. Observed model names include JPTP, MGTP, and APTDPT. Select the model matching the component before starting production.

Unauthorised model/program changes can apply sealant incorrectly and may create unexpected robot paths. Only authorised persons may change models or teach programs (OA-P7-011).""",
    )

    h2(doc, "6.5 Manual Mode")
    body(
        doc,
        """MANUAL mode is available on the HMI for setup and recovery tasks. Manual operations must be performed only by trained personnel with an understanding of robot motion and door behaviour.""",
    )
    alert_box(
        doc,
        "CAUTION",
        ["Manual mode does not remove the need for LOTO when hands enter the cell for tool or tip work."],
    )

    h2(doc, "6.6 Parking and Purging")
    body(
        doc,
        f"""The machine provides parking and purging cups for the dispenser tip. HMI Settings observed: Idle Time To Go Parking = {DATA['hmi']['settings_observed']['idle_time_to_go_parking_min']} Min; Parking & Purging Sequence = {DATA['hmi']['settings_observed']['parking_and_purging_sequence']}.

Purpose: if the tip is left idle, sealant can dry and block dispensing. Use automatic parking/purging and the robot DISPENSER PURGE control as trained.""",
    )
    alert_box(
        doc,
        "NOTICE",
        [
            "Confirm process idle limit with sealant SDS and OEM (OA-P7-006).",
            "Image placeholders: Images/dispenser parking.jpg; dispenser purging.jpg; Dispenser.jpg.",
        ],
    )

    flowchart(
        doc,
        "6.7 Shutdown",
        [
            "START",
            "Complete or abort cycle safely",
            "Park or purge dispenser tip",
            "Clear HMI to idle state",
            "Turn isolator OFF if end of shift",
            "Isolate air if required by site",
            "Clean work area",
            "Record issues in log",
            "END",
        ],
    )

    h2(doc, "6.8 Forbidden Actions")
    bullet(
        doc,
        [
            "Bypass two-hand, door sensors, or E-Stops",
            "Operate with side doors open in production",
            "Disable Front Automatic Safety Door or Side Door Safety Sensor settings without formal change control",
            "Teach robot programs without authorisation",
            "Clear tip blockage with unguarded hands while power/air available",
            "Work in the panel without electrical LOTO",
        ],
    )


def section_7(doc):
    h1(doc, "7. Maintenance & Care")
    h2(doc, "7.1 General Rules")
    body(
        doc,
        """All maintenance inside the process zone or panel requires LOTO (Section 4). E-Stop alone is not isolation.

Only skilled persons may open the control panel, modify pneumatics, or change electrical components.""",
    )

    h2(doc, "7.2 Recommended Checks")
    body(doc, "Daily (operator / end of shift):")
    bullet(
        doc,
        [
            "Visual check of doors, two-hand buttons, E-Stops, tip condition",
            "Confirm no damaged cables or air leaks",
            "Empty purge waste as required; keep parking/purge cups clean",
            "Confirm HMI free of standing alarms before leaving",
        ],
    )
    body(doc, "Weekly / periodic (skilled):")
    bullet(
        doc,
        [
            "FRL drain and filter condition (SMC unit)",
            "Omron side-door switch alignment and mounting",
            "Front door free travel and pneumatic fittings",
            "Isolator and panel door condition",
            "Robot mechanical play / cable track condition (per Fisnar OEM guidance)",
        ],
    )
    alert_box(
        doc,
        "PLACEHOLDER",
        ["[PLACEHOLDER — OEM PM schedule and Fisnar service intervals — attach when available]"],
    )

    h2(doc, "7.3 Sealant Tip and Dispenser Care")
    bullet(
        doc,
        [
            "Follow sealant manufacturer instructions for storage and pot life.",
            "Use parking/purge to reduce tip clogging.",
            "Replace tips/cartridges only after energy isolation if hands enter the motion zone.",
            "Dispose of waste sealant per site chemical rules and SDS.",
        ],
    )

    h2(doc, "7.4 After Maintenance")
    bullet(
        doc,
        [
            "Restore all guards and close doors",
            "Remove locks only after tools and persons are clear",
            "Function-test E-Stops, two-hand start, door behaviour, and air-pressure alarms before production release",
            "Record work in Appendix A log",
        ],
    )


def section_8(doc):
    h1(doc, "8. HMI, Alarms & Controls")
    h2(doc, "8.1 HMI Overview")
    body(
        doc,
        """Mitsubishi Electric HMI application title: GLUE DISPENSER.

Main menu functions observed: AUTO, MANUAL, MODEL SELECTION, SETTINGS, ALARM, ALARM HISTORY, DATA LOG, INPUT STATUS, OUTPUT STATUS.

Process screen shows mode and door/air/robot status lamps, model, robot program number, OK/NG/TOTAL counts, cycle time, and dispenser pressure.""",
    )
    net = DATA["hmi"]["network_labels_on_screen"]
    alert_box(
        doc,
        "NOTICE",
        [
            f"Network labels on menu photo: PLC {net['plc_ip_label']}, HMI {net['hmi_ip_label']}.",
            "Confirm before network changes (OA-P7-009).",
        ],
    )

    h2(doc, "8.2 Settings Screen (as photographed)")
    s = DATA["hmi"]["settings_observed"]
    styled_table(
        doc,
        ["Parameter", "Value"],
        [
            ["Dispenser Inlet Pressure Min/Max", f"{s['dispenser_inlet_pressure_min_bar']} / {s['dispenser_inlet_pressure_max_bar']} Bar"],
            ["Pressure Check Delay", f"{s['pressure_check_delay_sec']} Sec"],
            ["Dispenser Pressure Check", s["dispenser_pressure_check"]],
            ["Side Door Safety Sensor", s["side_door_safety_sensor"]],
            ["Front Automatic Safety Door", s["front_automatic_safety_door"]],
            ["Idle Time To Go Parking", f"{s['idle_time_to_go_parking_min']} Min"],
            ["Parking & Purging Sequence", s["parking_and_purging_sequence"]],
        ],
        [7, 9],
    )
    alert_box(
        doc,
        "WARNING",
        [
            "Do not turn safety-related settings OFF (side door sensor, front automatic safety door) for production convenience.",
            "Any change requires documented authorisation and risk review.",
        ],
    )

    h2(doc, "8.3 Alarms Observed")
    styled_table(
        doc,
        ["Message / Code", "Notes"],
        [[a["message"], a["source"]] for a in DATA["alarms_observed"]],
        [7, 9],
    )
    alert_box(
        doc,
        "PLACEHOLDER",
        ["[PLACEHOLDER — Full HMI/PLC alarm catalogue — OA-P7-008]"],
    )

    flowchart(
        doc,
        "8.4 Alarm Response",
        [
            "START",
            "Stop and keep clear",
            "Read HMI live alarm",
            "Identify cause safely",
            "Apply LOTO if entry needed",
            "Correct root cause",
            "Reset E-Stop if used",
            "Clear alarm on HMI",
            "Test then resume",
            "END",
        ],
    )

    h2(doc, "8.5 Operator Hardware Controls")
    body(doc, DATA["operator_controls"]["notes"])
    bullet(
        doc,
        [
            "Left station: RESET (yellow), TWO HAND CYCLE START (green)",
            "Right station: TWO HAND CYCLE START (green), E STOP (red)",
            "Robot face: DISPENSER PURGE, PROG NO., RUN/TEACH, START, EMERGENCY STOP",
            "Panel: POWER ON/OFF isolator; SMC FRL",
        ],
    )


def section_9(doc):
    h1(doc, "9. Safety Features, Schematics & Compliance")
    h2(doc, "9.1 Safety Features Summary")
    rows = [[sd["id"], sd["name"], sd["detail"][:120] + ("…" if len(sd["detail"]) > 120 else "")] for sd in DATA["safety_devices"]]
    styled_table(doc, ["ID", "Device", "Detail (summary)"], rows, [2, 4.5, 9.5])
    alert_box(
        doc,
        "OPEN ACTION",
        [
            "OA-P7-005: Validate front automatic door performance and residual risk of no light curtain at the load face.",
            "PLr / category of safety functions are not claimed without a Machine Safety Assessment (OA-P7-001).",
        ],
    )

    h2(doc, "9.2 Electrical Schematic (placeholder for customer insert)")
    alert_box(
        doc,
        "PLACEHOLDER",
        [
            "[PLACEHOLDER — Electrical schematic drawing to be inserted by OEM/customer — OA-P7-003]",
            "When available, insert: single-line supply, main isolator, PLC I/O power, E-Stop circuit overview, panel layout.",
            "Photo reference for as-built panel interior: Images/PLC panel open.jpg (not a substitute for a controlled schematic).",
        ],
    )

    h2(doc, "9.3 Pneumatic Schematic (placeholder for customer insert)")
    alert_box(
        doc,
        "PLACEHOLDER",
        [
            "[PLACEHOLDER — Pneumatic schematic drawing to be inserted by OEM/customer — OA-P7-004]",
            "When available, insert: supply, FRL, isolation valve, front door actuators, dispenser air path.",
            "Photo references: Images/Pneumatic FRL.jpg; Front door pneumatic view.jpg.",
        ],
    )

    h2(doc, "9.4 Applicable Standards (reference)")
    body(
        doc,
        """The following standards and regulations are referenced as design/instruction principles for this manual. Clause numbers are not fabricated.""",
    )
    bullet(doc, DATA["standards_in_use"])
    body(doc, "Explicitly not referenced in this manual: " + "; ".join(DATA["standards_excluded"]) + ".")

    h2(doc, "9.5 Liability Disclaimer")
    body(
        doc,
        """This manual is prepared from available photographs and intake information for the Robotic Sealant Dispenser Machine. It does not replace a formal Machine Safety Assessment, OEM electrical/pneumatic design packages, or site-specific safe work procedures.

Rani Enterprises (OEM), the end user, and operators remain responsible for correct installation, training, maintenance, and safe use. The Certified Safety Professional preparing this draft is not liable for use of incomplete data, defeated safeguards, unauthorised modifications, or failure to close Open Actions before production release.

Where information was unavailable, PLACEHOLDER text and Open Actions are used. Closing those actions is a condition of treating this document as complete.""",
    )

    h2(doc, "9.6 Document Sign-Off")
    body(
        doc,
        """Prepared by: Certified Safety Professional
Role: Independent Safety Consultant (instruction manual authoring)

Reviewed by: ___________________________  Date: ___________

Approved by (OEM): _____________________  Date: ___________

Approved by (End User): _________________  Date: ___________

Sign-off appears once in this manual (this subsection only).""",
    )


def appendix_a(doc):
    h1(doc, "Appendix A — Maintenance Log Template")
    body(doc, "Use this log for routine and corrective maintenance. Attach additional pages as needed.")
    styled_table(
        doc,
        ["Date", "Task performed", "Parts replaced", "Technician", "LOTO (Y/N)", "Next due", "Remarks"],
        [["", "", "", "", "", "", ""] for _ in range(6)],
        [2, 3.5, 2.5, 2.5, 1.5, 2, 2.5],
    )


def appendix_b(doc):
    h1(doc, "Appendix B — LOTO Log Template")
    body(doc, "Use this log whenever Lockout/Tagout is applied.")
    styled_table(
        doc,
        ["Date", "Time on", "Time off", "Zone", "LP-01", "LP-02", "Worker", "Lock ID", "Reason", "Zero energy", "Supervisor"],
        [["", "", "", "", "", "", "", "", "", "", ""] for _ in range(5)],
        [1.5, 1.3, 1.3, 1.8, 1.2, 1.2, 1.8, 1.3, 2, 1.5, 1.8],
    )


def main():
    doc = Document()
    build_header_footer(doc)
    cover(doc)
    revision_page(doc)
    toc_page(doc)
    section_1(doc)
    section_2(doc)
    section_3(doc)
    section_4(doc)
    section_5(doc)
    section_6(doc)
    section_7(doc)
    section_8(doc)
    section_9(doc)
    appendix_a(doc)
    appendix_b(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
