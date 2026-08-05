#!/usr/bin/env python3
"""Build Open Actions Register xlsx from machine_data.json (P7 Rani Enterprises)."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / "inputs" / "machine_data.json").read_text(encoding="utf-8"))
OUT = ROOT / "outputs" / "RaniEnterprises_RoboticSealantDispenser_OpenActions_V1.0.xlsx"

NAVY = "1F3864"
ALT = "D9E1F2"
HEADERS = ["ID", "Description", "Category", "Priority", "Responsibility", "Target Date", "Status"]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Open Actions"

    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=NAVY)
    body_font = Font(name="Calibri", size=10)
    alt_fill = PatternFill("solid", fgColor=ALT)
    thin = Border(
        left=Side(style="thin", color=NAVY),
        right=Side(style="thin", color=NAVY),
        top=Side(style="thin", color=NAVY),
        bottom=Side(style="thin", color=NAVY),
    )
    wrap = Alignment(wrap_text=True, vertical="center")

    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin

    for ri, oa in enumerate(DATA["open_actions"], 2):
        values = [
            oa["id"],
            oa["description"],
            oa["category"],
            oa["priority"],
            oa["responsibility"],
            oa["target_date"],
            oa["status"],
        ]
        for ci, val in enumerate(values, 1):
            cell = ws.cell(ri, ci, val)
            cell.font = body_font
            cell.alignment = wrap
            cell.border = thin
            if ri % 2 == 0:
                cell.fill = alt_fill

    widths = [12, 70, 16, 12, 28, 12, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 22
    for r in range(2, 2 + len(DATA["open_actions"])):
        ws.row_dimensions[r].height = 48

    ws2 = wb.create_sheet("Meta")
    ws2["A1"] = "Project"
    ws2["B1"] = DATA["project_code"]
    ws2["A2"] = "Machine"
    ws2["B2"] = DATA["machine_name"]
    ws2["A3"] = "Manufacturer"
    ws2["B3"] = DATA["manufacturer"]
    ws2["A4"] = "Status"
    ws2["B4"] = DATA["status"]
    ws2["A5"] = "Revision"
    ws2["B5"] = DATA["revision"]
    ws2["A6"] = "Note"
    ws2["B6"] = "Standalone Open Actions Register — not embedded in OSM"
    for row in ws2.iter_rows(min_row=1, max_row=6, max_col=2):
        for cell in row:
            cell.font = Font(name="Calibri", size=11)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
