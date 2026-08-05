import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def build_open_actions_excel():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    inputs_dir = os.path.join(base_dir, "inputs")
    outputs_dir = os.path.join(base_dir, "outputs")
    os.makedirs(outputs_dir, exist_ok=True)

    with open(os.path.join(inputs_dir, "machine_data.json"), "r", encoding="utf-8") as f:
        data = json.load(f)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Open Actions Register"

    # Styling definitions
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="000000")
    alt_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='1F3864'),
        right=Side(style='thin', color='1F3864'),
        top=Side(style='thin', color='1F3864'),
        bottom=Side(style='thin', color='1F3864')
    )

    # Title Block
    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = f"OPEN ACTIONS REGISTER — {data['client']} ({data['machine_name']})"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1F3864")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    headers = ["Action ID", "Description of Gap / Safety Action", "Category", "Priority", "Responsibility", "Target Date", "Status"]
    
    ws.append([]) # Empty row
    ws.append(headers)

    # Style Header Row
    for col_num in range(1, 8):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Data Rows
    row_idx = 4
    for action in data.get("open_actions", []):
        row_values = [
            action.get("id", ""),
            action.get("description", ""),
            action.get("category", ""),
            action.get("priority", ""),
            action.get("responsibility", ""),
            action.get("target_date", ""),
            action.get("status", "")
        ]
        ws.append(row_values)
        
        # Style Data Row
        for col_num in range(1, 8):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if row_idx % 2 == 1:
                cell.fill = alt_fill

        row_idx += 1

    # Adjust Column Widths
    col_widths = [15, 45, 18, 12, 28, 20, 12]
    for idx, width in enumerate(col_widths, start=1):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws.column_dimensions[col_letter].width = width

    output_path = os.path.join(outputs_dir, "RaniEnterprises_RoboticSealantDispenser_OpenActions_V1.0.xlsx")
    wb.save(output_path)
    print(f"Successfully generated Open Actions Register: {output_path}")

if __name__ == "__main__":
    build_open_actions_excel()
